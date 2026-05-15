import os
import io
import json
import time
import urllib.request
import openpyxl
from datetime import datetime
from flask import (Flask, render_template, request, redirect,
                   url_for, flash, jsonify, send_file)

import proof_engine

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'prodough-proof-site-2024-local')
app.config['MAX_CONTENT_LENGTH'] = 500 * 1024 * 1024  # 500 MB

BASE_DIR              = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR            = os.path.join(BASE_DIR, 'uploads')
GTIN_STORE_PATH       = os.path.join(UPLOAD_DIR, 'gtin_store.json')
GTIN_SHEET_CFG_PATH   = os.path.join(UPLOAD_DIR, 'gtin_sheet_config.json')
os.makedirs(UPLOAD_DIR, exist_ok=True)

_sheet_cache: dict = {'rows': None, 'fetched_at': 0.0, 'url': ''}
_SHEET_CACHE_TTL = 300  # seconds

ALLOWED_EXTS = {'.pdf', '.ai', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.eps'}

CHECK_LABELS = {
    'gtin':     'GTIN / Barcode',
    'nfp':      'Front Call-Out vs NFP',
    'eyemark':  'Eyemark Contrast',
    'spelling': 'Spelling / Brand Name',
    'fda':      'FDA Audit Risk',
}


@app.template_filter('basename_filter')
def basename_filter(path):
    return os.path.basename(path) if path else ''


# ── GTIN store ────────────────────────────────────────────────────────────────────

def _load_gtin_store():
    if os.path.exists(GTIN_STORE_PATH):
        try:
            with open(GTIN_STORE_PATH) as f:
                return json.load(f)
        except Exception:
            pass
    return None


def _save_gtin_store(rows: list, source_filename: str) -> dict:
    for row in rows:
        gtin = str(row.get('gtin', ''))
        if not gtin:
            row['_valid'] = False
            row['_error'] = 'Missing GTIN'
        elif not gtin.isdigit():
            row['_valid'] = False
            row['_error'] = 'Non-numeric characters'
        elif len(gtin) != 12:
            row['_valid'] = False
            row['_error'] = f'{len(gtin)} digits (expect 12)'
        else:
            row['_valid'] = True
            row['_error'] = None
    store = {
        'uploaded_at': datetime.now().isoformat(),
        'filename': source_filename,
        'row_count': len(rows),
        'error_count': sum(1 for r in rows if not r.get('_valid')),
        'rows': rows,
    }
    with open(GTIN_STORE_PATH, 'w') as f:
        json.dump(store, f)
    return store


# ── Google Sheets sync ──────────────────────────────────────────────────────

def _load_sheet_config() -> dict:
    if os.path.exists(GTIN_SHEET_CFG_PATH):
        try:
            with open(GTIN_SHEET_CFG_PATH) as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def _save_sheet_config(cfg: dict):
    with open(GTIN_SHEET_CFG_PATH, 'w') as f:
        json.dump(cfg, f)


def _sheet_url_to_csv(url: str) -> str:
    """Convert any Google Sheets share/edit URL to a direct CSV export URL."""
    import re
    if 'export?format=csv' in url or 'output=csv' in url:
        return url
    m = re.search(r'/spreadsheets/d/([a-zA-Z0-9_-]+)', url)
    if not m:
        raise ValueError('Not a recognized Google Sheets URL. Paste the URL from your browser address bar.')
    sheet_id = m.group(1)
    gid_m = re.search(r'[#&?]gid=(\d+)', url)
    gid_part = f'&gid={gid_m.group(1)}' if gid_m else ''
    return f'https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv{gid_part}'


def _fetch_sheet_rows(csv_url: str) -> list:
    """Download and parse a published Google Sheet as CSV."""
    import csv as _csv
    req = urllib.request.Request(csv_url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=15) as resp:
        content = resp.read().decode('utf-8-sig')
    reader = _csv.DictReader(io.StringIO(content))
    rows = []
    for row in reader:
        d = {k.strip().lower(): (v.strip() if v else '') for k, v in row.items() if k}
        gtin_raw = d.get('gtin/barcode#') or d.get('gtin') or d.get('barcode') or d.get('upc') or ''
        flavor   = d.get('flavor') or d.get('product name') or d.get('name') or ''
        sku      = d.get('sku') or ''
        gtin_str = str(gtin_raw).strip().replace(' ', '')
        if '.' in gtin_str and gtin_str.replace('.', '').isdigit():
            try:
                gtin_str = str(int(float(gtin_str)))
            except (ValueError, OverflowError):
                pass
        if gtin_str and gtin_str not in ('', 'nan'):
            rows.append({'gtin': gtin_str, 'flavor': str(flavor).strip().lower(), 'sku': str(sku).strip()})
    return rows


def _get_sheet_gtin_rows(force: bool = False) -> list:
    """Return GTIN rows from the synced Google Sheet, with a 5-minute in-memory cache."""
    global _sheet_cache
    cfg = _load_sheet_config()
    url = cfg.get('sheet_url', '')
    if not url:
        return []
    now = time.time()
    if (not force and _sheet_cache['url'] == url
            and _sheet_cache['rows'] is not None
            and now - _sheet_cache['fetched_at'] < _SHEET_CACHE_TTL):
        return _sheet_cache['rows']
    try:
        csv_url = _sheet_url_to_csv(url)
        rows = _fetch_sheet_rows(csv_url)
        _sheet_cache = {'rows': rows, 'fetched_at': now, 'url': url}
        cfg['last_synced'] = datetime.now().isoformat()
        cfg['row_count'] = len(rows)
        _save_sheet_config(cfg)
        return rows
    except Exception:
        return _sheet_cache['rows'] if _sheet_cache['rows'] is not None else []


# ── Pages ───────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    return render_template('proof.html', jobs=proof_engine.list_jobs())


@app.route('/gtin')
def gtin_page():
    store = _load_gtin_store()
    sheet_cfg = _load_sheet_config()
    # If a sheet is configured, show the live sheet rows in the table
    if sheet_cfg.get('sheet_url') and not store:
        rows = _get_sheet_gtin_rows()
        if rows:
            store = _save_gtin_store(rows, 'Google Sheets (auto-synced)')
    return render_template('gtin.html', store=store, sheet_cfg=sheet_cfg)


@app.route('/upload', methods=['POST'])
def upload():
    artwork_files = request.files.getlist('artwork')
    gtin_file     = request.files.get('gtin_list')

    if not artwork_files or all(f.filename == '' for f in artwork_files):
        flash('Please select at least one artwork file.', 'danger')
        return redirect(url_for('index'))

    # GTIN data: uploaded file takes priority; fall back to synced Google Sheet
    gtin_rows = []
    if gtin_file and gtin_file.filename:
        try:
            gtin_rows = _parse_gtin_list(gtin_file.read(), gtin_file.filename)
            if gtin_rows:
                _save_gtin_store(gtin_rows, gtin_file.filename)
            else:
                flash('The uploaded GTIN file appears empty — falling back to synced sheet data.', 'warning')
        except Exception as exc:
            flash(f'Could not read the uploaded GTIN file ({exc}) — falling back to synced sheet data.', 'warning')

    if not gtin_rows:
        gtin_rows = _get_sheet_gtin_rows()

    if not gtin_rows:
        flash(
            'No GTIN data available. Either upload a GTIN list on this page, or connect your '
            'Google Sheet on the GTIN List page so the barcode check has a master list to match against.',
            'warning',
        )

    job_id  = proof_engine.create_job([f.filename for f in artwork_files if f.filename])
    job_dir = os.path.join(UPLOAD_DIR, job_id)
    os.makedirs(job_dir, exist_ok=True)

    saved, skipped = [], []
    for f in artwork_files:
        if not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ALLOWED_EXTS:
            skipped.append(f.filename)
            continue
        safe = _safe_name(f.filename)
        dest = os.path.join(job_dir, safe)
        f.save(dest)
        saved.append(dest)

    if skipped:
        flash(f'Skipped unsupported file(s): {", ".join(skipped)}', 'warning')
    if not saved:
        flash('No supported artwork files were uploaded.', 'danger')
        return redirect(url_for('index'))

    proof_engine.start_job(job_id, saved, gtin_rows, job_dir)
    return redirect(url_for('result', job_id=job_id))


@app.route('/result/<job_id>')
def result(job_id):
    job = proof_engine.get_job(job_id)
    if not job:
        flash('Job not found.', 'danger')
        return redirect(url_for('index'))
    return render_template('result.html', job=job, job_id=job_id)


@app.route('/summary/<job_id>')
def summary(job_id):
    job = proof_engine.get_job(job_id)
    if not job:
        flash('Job not found.', 'danger')
        return redirect(url_for('index'))
    if job['status'] != 'done':
        return redirect(url_for('result', job_id=job_id))
    return render_template('summary.html', job=job, job_id=job_id)


@app.route('/history')
def history():
    jobs = proof_engine.list_jobs()
    return render_template('history.html', jobs=jobs)


@app.route('/viewer/<job_id>')
def viewer(job_id):
    job = proof_engine.get_job(job_id)
    if not job:
        flash('Job not found.', 'danger')
        return redirect(url_for('index'))

    verify_map = {}
    for r in job.get('results', []):
        items = []
        for check_name, check in r.get('checks', {}).items():
            for issue in check.get('issues', []):
                msg = issue.get('message', '')
                if 'verify' in msg.lower() or 'manually' in msg.lower():
                    items.append({
                        'check_label': CHECK_LABELS.get(check_name, check_name),
                        'severity': issue['severity'],
                        'message': msg,
                    })
            for note in check.get('notes', []):
                if 'verify' in note.lower() or 'manually' in note.lower():
                    items.append({
                        'check_label': CHECK_LABELS.get(check_name, check_name),
                        'severity': 'info',
                        'message': note,
                    })
        verify_map[r['filename']] = items

    return render_template('viewer.html', job=job, job_id=job_id, verify_map=verify_map)


# ── API ───────────────────────────────────────────────────────────────────────────

@app.route('/api/status/<job_id>')
def api_status(job_id):
    job = proof_engine.get_job(job_id)
    if not job:
        return jsonify({'error': 'not found'}), 404
    return jsonify({
        'status':       job['status'],
        'progress':     job['progress'],
        'current_file': job['current_file'],
        'error':        job['error'],
    })


@app.route('/api/result/<job_id>')
def api_result(job_id):
    job = proof_engine.get_job(job_id)
    if not job:
        return jsonify({'error': 'not found'}), 404
    return jsonify(job)


@app.route('/api/gtin-sheet', methods=['POST'])
def api_gtin_sheet_save():
    data = request.get_json() or {}
    url = data.get('url', '').strip()
    if url:
        try:
            _sheet_url_to_csv(url)  # validate it's a Sheets URL
        except ValueError as e:
            return jsonify({'ok': False, 'error': str(e)}), 400
    cfg = _load_sheet_config()
    cfg['sheet_url'] = url
    cfg.pop('last_synced', None)
    cfg.pop('row_count', None)
    _save_sheet_config(cfg)
    global _sheet_cache
    _sheet_cache = {'rows': None, 'fetched_at': 0.0, 'url': ''}
    return jsonify({'ok': True})


@app.route('/api/gtin-sheet/sync', methods=['POST'])
def api_gtin_sheet_sync():
    cfg = _load_sheet_config()
    if not cfg.get('sheet_url'):
        return jsonify({'ok': False, 'error': 'No sheet URL configured'}), 400
    try:
        rows = _get_sheet_gtin_rows(force=True)
        if rows:
            _save_gtin_store(rows, 'Google Sheets (auto-synced)')
        cfg = _load_sheet_config()
        return jsonify({'ok': True, 'row_count': len(rows), 'last_synced': cfg.get('last_synced', '')})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


@app.route('/api/dismiss/<job_id>', methods=['POST'])
def api_dismiss(job_id):
    data = request.get_json()
    if not data:
        return jsonify({'error': 'no data'}), 400
    ok = proof_engine.set_dismissal(
        job_id,
        data.get('filename', ''),
        data.get('check', ''),
        int(data.get('index', 0)),
        bool(data.get('dismissed', True)),
    )
    if not ok:
        return jsonify({'error': 'job not found'}), 404
    return jsonify({'ok': True})


@app.route('/report/<job_id>')
def download_report(job_id):
    job = proof_engine.get_job(job_id)
    if not job or job['status'] != 'done':
        flash('Report not ready — job must complete first.', 'danger')
        return redirect(url_for('index'))
    buf = _generate_report(job)
    fname = f'ProDough_Proof_{job_id}.xlsx'
    return send_file(
        buf, download_name=fname, as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/image/<job_id>/<path:filename>')
def serve_image(job_id, filename):
    job_dir  = os.path.realpath(os.path.join(UPLOAD_DIR, job_id))
    img_path = os.path.realpath(os.path.join(job_dir, filename))
    if not img_path.startswith(job_dir):
        return 'Forbidden', 403
    if not os.path.exists(img_path):
        return 'Not found', 404
    return send_file(img_path, mimetype='image/png')


# ── Helpers ─────────────────────────────────────────────────────────────────────

def _safe_name(name: str) -> str:
    keep = set('abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789._- ')
    return ''.join(c if c in keep else '_' for c in name)


def _parse_gtin_list(data: bytes, filename: str) -> list:
    rows = []
    if filename.lower().endswith('.csv'):
        import csv
        reader = csv.DictReader(io.TextIOWrapper(io.BytesIO(data), errors='replace'))
        for row in reader:
            rows.append({k.strip().lower(): v.strip() for k, v in row.items()})
    else:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            headers = None
            for row in ws.iter_rows(values_only=True):
                if headers is None:
                    if row and any(row):
                        headers = [str(c).strip().lower() if c else '' for c in row]
                    continue
                if not any(row):
                    continue
                d = dict(zip(headers, row))
                gtin_raw = d.get('gtin/barcode#') or d.get('gtin') or d.get('barcode') or d.get('upc') or ''
                flavor   = d.get('flavor') or d.get('product name') or d.get('name') or ''
                sku      = d.get('sku') or ''
                # Excel stores long numbers as floats (e.g. 850012345678.0) — convert via int first
                if isinstance(gtin_raw, (int, float)):
                    gtin_str = str(int(gtin_raw))
                else:
                    gtin_str = str(gtin_raw).strip().replace(' ', '')
                    if '.' in gtin_str and gtin_str.replace('.', '').replace('-', '').isdigit():
                        try:
                            gtin_str = str(int(float(gtin_str)))
                        except (ValueError, OverflowError):
                            pass
                rows.append({
                    'gtin':   gtin_str,
                    'flavor': str(flavor).strip().lower(),
                    'sku':    str(sku).strip(),
                })
    return [r for r in rows if r.get('gtin') and r['gtin'] not in ('', 'nan')]


def _generate_report(job: dict) -> io.BytesIO:
    from openpyxl.styles import Font, PatternFill, Alignment

    dismissals = job.get('dismissals', {})

    def is_dismissed(filename, check_name, idx):
        return dismissals.get(f"{filename}|{check_name}|{idx}", False)

    hdr_font  = Font(bold=True, color='FFFFFF')
    hdr_blue  = PatternFill('solid', fgColor='1A56A0')
    hdr_gray  = PatternFill('solid', fgColor='6C757D')
    fill_crit = PatternFill('solid', fgColor='FFD0D0')
    fill_warn = PatternFill('solid', fgColor='FFF3CD')
    fill_ok   = PatternFill('solid', fgColor='D4EDDA')
    wrap      = Alignment(wrap_text=True, vertical='top')
    center    = Alignment(horizontal='center', vertical='top')

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Summary'
    ws1.merge_cells('A1:E1')
    ws1['A1'] = 'ProDough Artwork Proof Report'
    ws1['A1'].font = Font(bold=True, size=14)
    ws1['A2'] = f'Job: {job["id"]}'
    ws1['B2'] = f'Date: {job["created"][:10]}'
    ws1['C2'] = f'Files proofed: {len(job["results"])}'
    ws1.append([])
    for col, hdr in enumerate(['File', 'Overall Status', 'Critical', 'Warnings', 'Notes'], 1):
        c = ws1.cell(row=4, column=col, value=hdr)
        c.font = hdr_font
        c.fill = hdr_blue
        c.alignment = center

    for r in job['results']:
        ws1.append([r['filename'], r.get('severity', 'error').upper(),
                    r.get('critical_count', 0), r.get('warning_count', 0), r.get('info_count', 0)])
        sev = r.get('severity', 'error')
        f = fill_crit if sev == 'critical' else fill_warn if sev == 'warning' else fill_ok if sev == 'clean' else None
        if f:
            for col in range(1, 6):
                ws1.cell(row=ws1.max_row, column=col).fill = f

    ws1.column_dimensions['A'].width = 45
    for col in ['B', 'C', 'D', 'E']:
        ws1.column_dimensions[col].width = 14

    # ── Sheet 2: Issues to Fix (non-dismissed) ─────────────────────────────────────────
    ws2 = wb.create_sheet('Issues to Fix')
    for col, hdr in enumerate(['File', 'Check', 'Severity', 'Issue / Required Action'], 1):
        c = ws2.cell(row=1, column=col, value=hdr)
        c.font = hdr_font
        c.fill = hdr_blue
        c.alignment = center

    row_idx = 2
    for r in job['results']:
        for check_name, check in r.get('checks', {}).items():
            for i, issue in enumerate(check.get('issues', [])):
                if is_dismissed(r['filename'], check_name, i):
                    continue
                sev = issue['severity']
                vals = [r['filename'], CHECK_LABELS.get(check_name, check_name),
                        sev.upper(), issue['message']]
                for col, val in enumerate(vals, 1):
                    c = ws2.cell(row=row_idx, column=col, value=val)
                    c.fill = fill_crit if sev == 'critical' else fill_warn if sev == 'warning' else PatternFill()
                    c.alignment = wrap
                row_idx += 1

    ws2.column_dimensions['A'].width = 40
    ws2.column_dimensions['B'].width = 22
    ws2.column_dimensions['C'].width = 12
    ws2.column_dimensions['D'].width = 80

    # ── Sheet 3: Reviewed / Dismissed ────────────────────────────────────────────────
    if dismissals:
        ws3 = wb.create_sheet('Reviewed OK')
        for col, hdr in enumerate(['File', 'Check', 'Severity', 'Issue (Reviewed — No Action Needed)'], 1):
            c = ws3.cell(row=1, column=col, value=hdr)
            c.font = hdr_font
            c.fill = hdr_gray
            c.alignment = center

        row_idx = 2
        for r in job['results']:
            for check_name, check in r.get('checks', {}).items():
                for i, issue in enumerate(check.get('issues', [])):
                    if not is_dismissed(r['filename'], check_name, i):
                        continue
                    vals = [r['filename'], CHECK_LABELS.get(check_name, check_name),
                            issue['severity'].upper(), issue['message']]
                    for col, val in enumerate(vals, 1):
                        ws3.cell(row=row_idx, column=col, value=val).alignment = wrap
                    row_idx += 1

        ws3.column_dimensions['A'].width = 40
        ws3.column_dimensions['B'].width = 22
        ws3.column_dimensions['C'].width = 12
        ws3.column_dimensions['D'].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
